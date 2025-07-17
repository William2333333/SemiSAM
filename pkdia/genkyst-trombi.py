#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import numpy as np
from exams.exam_genkyst import exam_genkyst
from skimage.transform import resize, rotate
from skimage.exposure import rescale_intensity
from skimage import io, img_as_ubyte
from openpyxl import load_workbook
from current import current_xlsx

modality = 'T2'
offset = 40
size = 256
folder = './../../../../Desktop/trombinoscope/'

def most_visible_slice(modality, mask):
    maxactive, maxyz = 0, 0
    if modality == 'T2':
        ys = mask.shape[1]
        for y in range(0,ys):
            activepx = len(np.where(mask.get_fdata()[:,y,:]>0.)[0])
            if activepx > maxactive:
                maxactive = activepx
                maxyz = y
    elif modality == 'CT':
        zs = mask.shape[2]
        for z in range(0,zs):
            activepx = len(np.where(mask.get_fdata()[:,:,z]>0.)[0])
            if activepx > maxactive:
                maxactive = activepx
                maxyz = z        
    return maxyz

def extraction(modality, img, mask, yz, offset):
    if modality == 'T2':
        (posX, posZ) = np.where(mask.get_fdata()[:,yz,:] > 0.)            
        xmin_, zmin_, xmax_, zmax_ = min(posX), min(posZ), max(posX), max(posZ)
        xmin = xmin_-offset if xmin_-offset>0 else 0
        xmax = xmax_+offset if xmax_+offset<img.get_fdata()[:,yz,:].shape[0] else img.get_fdata()[:,yz,:].shape[0]-1
        zmin = zmin_-offset if zmin_-offset>0 else 0
        zmax = zmax_+offset if zmax_+offset<img.get_fdata()[:,yz,:].shape[1] else img.get_fdata()[:,yz,:].shape[1]-1
        return img.get_fdata()[xmin:xmax,yz,zmin:zmax]
    elif modality == 'CT':
        (posX, posY) = np.where(mask.get_fdata()[:,:,yz] > 0.)            
        xmin_, ymin_, xmax_, ymax_ = min(posX), min(posY), max(posX), max(posY)
        xmin = xmin_-offset if xmin_-offset>0 else 0
        xmax = xmax_+offset if xmax_+offset<img.get_fdata()[:,:,yz].shape[0] else img.get_fdata()[:,:,yz].shape[0]-1
        ymin = ymin_-offset if ymin_-offset>0 else 0
        ymax = ymax_+offset if ymax_+offset<img.get_fdata()[:,:,yz].shape[1] else img.get_fdata()[:,:,yz].shape[1]-1
        return img.get_fdata()[xmin:xmax,ymin:ymax,yz]

wb = load_workbook(current_xlsx()[0])
sh = wb['main']
list_ids = [sh.cell(row=rownum, column=1).value for rownum in range(2,sh.max_row+1)]
list_series = [sh.cell(row=rownum, column=8).value for rownum in range(2,sh.max_row+1)]
print(list_ids)
print(list_series)

for idx, id_ in enumerate(list_ids):
    print(id_)
    id_ = int(list_ids[idx])
    serie = int(list_series[idx])
    exam = exam_genkyst(id_, serie, modality, False)
    
    if modality == 'T2':
        exist = exam.T2_exist
    elif modality == 'CT':
        exist = exam.CT_exist
    
    if exist:
        exam.exam_upload(modality)
        if exam.kid_annot:

            yzLK = most_visible_slice(modality, exam.LK)
            yzRK = most_visible_slice(modality, exam.RK)
            print(yzLK,yzRK)
            
            if modality == 'T2': 
                imgLK = extraction(modality, exam.T2, exam.LK, yzLK, offset)
                imgRK = extraction(modality, exam.T2, exam.RK, yzRK, offset)
            elif modality == 'CT':
                imgLK = extraction(modality, exam.CT, exam.LK, yzLK, offset)
                imgRK = extraction(modality, exam.CT, exam.RK, yzRK, offset)
            print(imgLK.shape, imgRK.shape)
            
            imgLK = rotate(resize(imgLK[::-1,:], output_shape=(size,size), preserve_range=True), 90, preserve_range=True)
            imgRK = rotate(resize(imgRK[::-1,:], output_shape=(size,size), preserve_range=True), 90, preserve_range=True)         
            
            mingreyscaleLK, maxgreyscaleLK = np.percentile(imgLK,(5,95)) 
            mingreyscaleRK, maxgreyscaleRK = np.percentile(imgRK,(5,95))
            mingreyscale = min(mingreyscaleLK, mingreyscaleRK)
            maxgreyscale = max(maxgreyscaleLK, maxgreyscaleRK)
            
            imgLK = rescale_intensity(imgLK, in_range=(mingreyscale,maxgreyscale), out_range=(0,1))
            imgRK = rescale_intensity(imgRK, in_range=(mingreyscale,maxgreyscale), out_range=(0,1))
            
            io.imsave(folder+exam.id+'-%0*d'%(2,exam.serie)+'-LK.png', img_as_ubyte(imgLK))
            io.imsave(folder+exam.id+'-%0*d'%(2,exam.serie)+'-RK.png', img_as_ubyte(imgRK))